import os
import json
import hashlib
import re
from datetime import datetime, date

import openpyxl

from app.storage.lancedb_store import get_spreadsheet_table


FILE_NAME = "IT_Direct_Hire_Companies_2026.xlsx"

BASE_DIR = os.path.dirname(os.path.abspath(__file__))
FILE_PATH = os.path.join(
    BASE_DIR,
    "watched_folder",
    FILE_NAME
)


def json_safe(value):
    """Convert Excel/Python values into JSON-safe values."""

    if value is None:
        return None

    if isinstance(value, (datetime, date)):
        return value.isoformat()

    if hasattr(value, "item"):
        try:
            return value.item()
        except Exception:
            pass

    return value


def file_hash(path):
    sha = hashlib.sha256()

    with open(path, "rb") as f:
        while True:
            block = f.read(1024 * 1024)

            if not block:
                break

            sha.update(block)

    return sha.hexdigest()


def normalize_header(value):
    if value is None:
        return ""

    return " ".join(str(value).strip().split()).lower()


def find_header_row_domain_agnostic(ws, max_scan=10):
    """
    Domain-agnostic header detection.
    Scans the first 10 rows for candidate header rows containing multiple concise text labels.
    Calculates a score based on header keyword matching, cell count, and label conciseness,
    excluding rows containing URLs or email addresses.
    """
    best_row = 1
    best_score = -100.0

    scan_limit = min(ws.max_row, max_scan)
    header_keywords = {
        'name', 'type', 'location', 'area', 'city', 'state', 'website', 'link', 'url',
        'email', 'role', 'roles', 'skill', 'skills', 'ctc', 'salary', 'package', 'lpa',
        'apply', 'connect', 'priority', 'total', 'count', 'number', 'startup', 'size',
        'product', 'analytics', 'gcc', 'section', 'content', 'tip', 'detail', 'code',
        's.no', 'no', 'id', 'category', 'status', 'date', 'price', 'cost', 'score',
        'rate', 'title', 'description', 'comment', 'notes', 'user', 'author'
    }

    for r in range(1, scan_limit + 1):
        row_vals = [ws.cell(row=r, column=c).value for c in range(1, ws.max_column + 1)]
        non_empty = [(c, str(v).strip()) for c, v in enumerate(row_vals, 1) if v is not None and str(v).strip() != '']
        if len(non_empty) < 2:
            continue

        labels = [val for _, val in non_empty]
        has_url_or_email = any('http://' in val.lower() or 'https://' in val.lower() or '@' in val for val in labels)
        if has_url_or_email:
            continue

        matched_kw_count = 0
        for val in labels:
            words = set(re.findall(r'\b[a-z0-9]+\b', val.lower()))
            if header_keywords.intersection(words):
                matched_kw_count += 1

        avg_len = sum(len(v) for v in labels) / max(len(labels), 1)
        concise_bonus = 5.0 if avg_len < 30 else 0.0

        score = (matched_kw_count * 10.0) + (len(labels) * 3.0) + concise_bonus + (10.0 / r)
        if score > best_score:
            best_score = score
            best_row = r

    return best_row


def extract_sheet(ws, source_file, source_path, file_hash_value):
    header_row = find_header_row_domain_agnostic(ws)

    if header_row is None:
        print(f"SKIP {ws.title}: header row not found")
        return []

    raw_headers = [
        ws.cell(row=header_row, column=col).value
        for col in range(1, ws.max_column + 1)
    ]

    headers = []
    for idx, header in enumerate(raw_headers):
        if header is None or str(header).strip() == "":
            headers.append(f"Unnamed: {idx}")
        else:
            headers.append(str(header).strip())

    records = []

    for excel_row in range(header_row + 1, ws.max_row + 1):
        values = [
            ws.cell(row=excel_row, column=col).value
            for col in range(1, len(headers) + 1)
        ]

        non_empty_vals = [v for v in values if v is not None and str(v).strip() != ""]
        if not non_empty_vals:
            continue

        # Skip section divider / legend rows (single non-empty cell in multi-column sheet)
        if len(headers) >= 3 and len(non_empty_vals) == 1:
            first_val = str(non_empty_vals[0]).strip()
            if first_val.startswith("──") or first_val.endswith("──") or "🟣" in first_val or "🟢" in first_val or "🟠" in first_val:
                continue

        row_data = {}
        for header, value in zip(headers, values):
            row_data[header] = json_safe(value)

        search_parts = [
            f"File: {source_file}",
            f"Sheet: {ws.title}",
            f"Row: {excel_row}",
        ]

        for header, value in row_data.items():
            if value is None:
                continue
            search_parts.append(f"{header}: {value}")

        records.append({
            "chunk_id": f"{file_hash_value[:16]}_{ws.title}_{excel_row}",
            "source_file": source_file,
            "source_path": source_path,
            "file_hash": file_hash_value,
            "sheet_name": ws.title,
            "row_number": excel_row,
            "row_data": json.dumps(row_data, ensure_ascii=False, default=str),
            "search_text": " | ".join(search_parts),
        })

    sample_dict = json.loads(records[0]["row_data"]) if records else {}

    print("\n--------------------------------------------------")
    print(f"SHEET NAME:         {ws.title}")
    print(f"DETECTED HEADER ROW: {header_row}")
    print(f"DETECTED HEADERS:   {headers}")
    print(f"NUMBER OF DATA ROWS:{len(records)}")
    print(f"SAMPLE ROW DATA:    {json.dumps(sample_dict, ensure_ascii=False)[:180]}")
    print("--------------------------------------------------")

    return records


def main():
    if not os.path.exists(FILE_PATH):
        raise FileNotFoundError(FILE_PATH)

    source_path = os.path.abspath(FILE_PATH)
    source_file = os.path.basename(source_path)

    print("==================================================")
    print("VALIDATING IN-MEMORY EXTRACTION BEFORE DATABASE EDIT")
    print("SOURCE:", source_path)
    print("==================================================")

    workbook_hash = file_hash(source_path)
    workbook = openpyxl.load_workbook(source_path, data_only=True)

    all_records = []

    for ws in workbook.worksheets:
        records = extract_sheet(ws, source_file, source_path, workbook_hash)
        all_records.extend(records)

    if not all_records:
        raise RuntimeError("Validation failed: No records extracted!")

    # Verify no Unnamed headers in city sheets
    city_sheets = ["Bangalore", "Pune", "Hyderabad", "Indore", "Ahmedabad"]
    for rec in all_records:
        if rec["sheet_name"] in city_sheets:
            row_dict = json.loads(rec["row_data"])
            unnamed_keys = [k for k in row_dict.keys() if k.startswith("Unnamed:")]
            if unnamed_keys:
                raise ValueError(f"Validation Error: Sheet {rec['sheet_name']} contains Unnamed headers: {unnamed_keys}")

    print("\n==================================================")
    print("VALIDATION PASSED SUCCESSFULLY")
    print(f"TOTAL CORRECTED RECORDS TO INSERT: {len(all_records)}")
    print("==================================================")

    table = get_spreadsheet_table()

    escaped_path = source_path.replace("'", "''")
    escaped_file = source_file.replace("'", "''")

    print("\nDeleting old IT_Direct_Hire_Companies_2026.xlsx rows from LanceDB...")

    table.delete(f"source_path = '{escaped_path}'")
    table.delete(f"source_file = '{escaped_file}'")

    print("Old rows deleted.")

    print(f"Inserting {len(all_records)} corrected structured rows...")
    table.add(all_records)

    print("\nDONE.")
    print("Current LanceDB spreadsheet_rows count:", table.count_rows())


if __name__ == "__main__":
    main()
