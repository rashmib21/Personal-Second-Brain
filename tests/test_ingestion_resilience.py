import os
import shutil
import sys
import tempfile
import zipfile
import tarfile
import gzip

# Add project root to sys.path
BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
sys.path.insert(0, BASE_DIR)

from app.extractors.router import extract_file
from app.embeddings.embedding_router import generate_embedding
from app.celery_app.tasks.process_file import route_file
from app.storage.lancedb_store import total_chunks, get_table

def run_tests():
    print("=" * 70)
    print("RUNNING PERSONAL SECOND BRAIN INGESTION & RESILIENCE TESTS")
    print("=" * 70)

    test_dir = tempfile.mkdtemp(prefix="second_brain_test_files_")

    try:
        # 1. Plain Text File (.txt)
        txt_path = os.path.join(test_dir, "sample.txt")
        with open(txt_path, "w", encoding="utf-8") as f:
            f.write("Personal Second Brain test document about machine learning architectures.")

        res_txt = extract_file(txt_path)
        print("\n[TEST 1] Plain Text (.txt):", res_txt.get("status"))
        assert res_txt["status"] == "SUCCESS"
        assert "machine learning" in res_txt["text"]

        # 2. Markdown File (.md)
        md_path = os.path.join(test_dir, "notes.md")
        with open(md_path, "w", encoding="utf-8") as f:
            f.write("# Meeting Notes\nDiscussion on data pipelines and storage.")

        res_md = extract_file(md_path)
        print("[TEST 2] Markdown (.md):", res_md.get("status"))
        assert res_md["status"] == "SUCCESS"

        # 3. CSV Spreadsheet (.csv)
        csv_path = os.path.join(test_dir, "data.csv")
        with open(csv_path, "w", encoding="utf-8") as f:
            f.write("Name,Age,Role\nAlice,30,Data Engineer\nBob,25,Analyst\n")

        res_csv = extract_file(csv_path)
        print("[TEST 3] CSV Spreadsheet (.csv):", res_csv.get("status"))
        assert res_csv["status"] == "SUCCESS"
        assert "Data Engineer" in res_csv["text"]

        # 4. Word Document (.docx)
        docx_path = os.path.join(test_dir, "report.docx")
        from docx import Document
        doc = Document()
        doc.add_paragraph("This is a Word document sample paragraph.")
        doc.save(docx_path)

        res_docx = extract_file(docx_path)
        print("[TEST 4] Word Document (.docx):", res_docx.get("status"))
        assert res_docx["status"] == "SUCCESS"
        assert "Word document" in res_docx["text"]

        # 5. PowerPoint Presentation (.pptx)
        pptx_path = os.path.join(test_dir, "slides.pptx")
        from pptx import Presentation
        prs = Presentation()
        slide = prs.slides.add_slide(prs.slide_layouts[0])
        slide.shapes.title.text = "Second Brain Presentation"
        prs.save(pptx_path)

        res_pptx = extract_file(pptx_path)
        print("[TEST 5] PowerPoint (.pptx):", res_pptx.get("status"))
        assert res_pptx["status"] == "SUCCESS"
        assert "Second Brain Presentation" in res_pptx["text"]

        # 6. ZIP Archive (.zip) containing supported files
        zip_path = os.path.join(test_dir, "archive.zip")
        with zipfile.ZipFile(zip_path, "w") as z:
            z.write(txt_path, arcname="inside.txt")
            z.write(csv_path, arcname="data.csv")

        res_zip = extract_file(zip_path)
        print("[TEST 6] ZIP Archive (.zip):", res_zip.get("status"))
        assert res_zip["status"] == "SUCCESS"
        assert "machine learning" in res_zip["text"]

        # 7. TAR.GZ Archive (.tar.gz)
        tgz_path = os.path.join(test_dir, "archive.tar.gz")
        with tarfile.open(tgz_path, "w:gz") as tar:
            tar.add(txt_path, arcname="inside_tar.txt")

        res_tgz = extract_file(tgz_path)
        print("[TEST 7] TAR.GZ Archive (.tar.gz):", res_tgz.get("status"))
        assert res_tgz["status"] == "SUCCESS"
        assert "machine learning" in res_tgz["text"]

        # 8. Standalone GZ file (.gz)
        gz_path = os.path.join(test_dir, "standalone.txt.gz")
        with open(txt_path, "rb") as f_in, gzip.open(gz_path, "wb") as f_out:
            shutil.copyfileobj(f_in, f_out)

        res_gz = extract_file(gz_path)
        print("[TEST 8] Standalone GZ (.gz):", res_gz.get("status"))
        assert res_gz["status"] == "SUCCESS"

        # 9. Unsupported Source Code (.py)
        py_path = os.path.join(test_dir, "script.py")
        with open(py_path, "w", encoding="utf-8") as f:
            f.write("print('Hello World')")

        res_py = extract_file(py_path)
        print("[TEST 9] Source Code (.py) Graceful Skip:", res_py.get("status"))
        assert res_py["status"] == "UNSUPPORTED_FORMAT"
        assert res_py["text"] == ""

        # 10. Unsupported Legacy Presentation (.ppt)
        ppt_path = os.path.join(test_dir, "old.ppt")
        with open(ppt_path, "wb") as f:
            f.write(b"Dummy binary content")

        res_ppt = extract_file(ppt_path)
        print("[TEST 10] Legacy Binary PPT (.ppt) Graceful Skip:", res_ppt.get("status"))
        assert res_ppt["status"] == "UNSUPPORTED_FORMAT"

        # 11. Unsupported Archive Format (.rar / .7z)
        rar_path = os.path.join(test_dir, "archive.rar")
        with open(rar_path, "wb") as f:
            f.write(b"Rar!")

        res_rar = extract_file(rar_path)
        print("[TEST 11] Unsupported Archive (.rar) Graceful Skip:", res_rar.get("status"))
        assert res_rar["status"] == "UNSUPPORTED_FORMAT"

        # 12. Corrupt PDF Handling
        corrupt_pdf = os.path.join(test_dir, "corrupt.pdf")
        with open(corrupt_pdf, "wb") as f:
            f.write(b"%PDF-1.4 corrupt content")

        res_pdf = extract_file(corrupt_pdf)
        print("[TEST 12] Corrupt PDF (.pdf) Graceful Catch:", res_pdf.get("status"))
        assert res_pdf["status"] in ["CORRUPT_FILE", "NO_TEXT_EXTRACTED"]

        # 13. Embedding Dimension Check (Must be exactly 384 dims)
        print("\n[TEST 13] Embedding Vector Dimension Inspection:")
        sample_chunks = [
            ("text", "Sample chunk for document embedding"),
            ("image", "OCR text extracted from screenshot"),
            ("audio", "ASR audio transcript text"),
            ("video", "Video transcript text content"),
            ("presentation", "PowerPoint slide header and content"),
            ("archive", "Extracted content from compressed archive")
        ]

        for file_type, chunk in sample_chunks:
            vec = generate_embedding(file_type, chunk)
            dim = len(vec)
            print(f"  - file_type='{file_type}': dimension={dim}")
            assert dim == 384, f"Expected 384 dimensions, got {dim}"

        # 14. Task Execution & LanceDB Insertion
        print("\n[TEST 14] End-to-End Celery Task Processing for Sample Document:")
        event = {
            "path": txt_path,
            "file_type": "text",
            "file_hash": "test_hash_unique_12345"
        }
        res_task = route_file(event)
        print("  Task Result:", res_task)
        assert res_task["status"] == "processed"
        assert res_task["total_chunks"] >= 1

        print("\n" + "=" * 70)
        print("ALL INGESTION & RESILIENCE TESTS PASSED SUCCESSFULLY!")
        print("=" * 70)

    finally:
        shutil.rmtree(test_dir)

if __name__ == "__main__":
    run_tests()
