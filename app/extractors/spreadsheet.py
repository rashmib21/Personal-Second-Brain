import csv
import logging
import os
from openpyxl import load_workbook

logger = logging.getLogger(__name__)

def extract_spreadsheet(path):
    ext = os.path.splitext(path)[1].lower()

    # CSV file handling
    if ext == ".csv":
        try:
            rows = []
            with open(path, "r", encoding="utf-8", errors="replace") as f:
                reader = csv.reader(f)
                for row in reader:
                    line = " ".join(val.strip() for val in row if val.strip())
                    if line:
                        rows.append(line)
            full_text = "\n".join(rows)
            if not full_text.strip():
                return {"text": "", "status": "NO_TEXT_EXTRACTED", "error": None}
            return {"text": full_text, "status": "SUCCESS", "error": None}
        except Exception as e:
            logger.error(f"Error reading CSV {path}: {e}")
            return {"text": "", "status": "CORRUPT_FILE", "error": str(e)}

    # Excel / Spreadsheet handling
    if ext in [".xls", ".ods"]:
        logger.warning(f"Legacy/OpenDocument spreadsheet format ({ext}) parsing attempted: {path}")

    try:
        workbook = load_workbook(path, data_only=True)
        text_lines = []
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows(values_only=True):
                row_vals = [str(cell).strip() for cell in row if cell is not None and str(cell).strip() != ""]
                if row_vals:
                    text_lines.append(" ".join(row_vals))

        full_text = "\n".join(text_lines)
        if not full_text.strip():
            return {"text": "", "status": "NO_TEXT_EXTRACTED", "error": None}

        return {"text": full_text, "status": "SUCCESS", "error": None}

    except Exception as e:
        logger.error(f"Failed to extract spreadsheet {path}: {e}")
        status = "UNSUPPORTED_FORMAT" if ext in [".xls", ".ods"] else "CORRUPT_FILE"
        return {"text": "", "status": status, "error": str(e)}