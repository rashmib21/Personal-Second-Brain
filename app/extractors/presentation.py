import logging
import os

logger = logging.getLogger(__name__)

def extract_presentation(path):
    """
    Extracts text from PowerPoint (.pptx) files.
    Returns structured dict with status.
    """
    ext = os.path.splitext(path)[1].lower()
    if ext == ".ppt":
        logger.warning(f"Legacy binary presentation format (.ppt) not supported natively: {path}")
        return {
            "text": "",
            "status": "UNSUPPORTED_FORMAT",
            "error": "Legacy binary .ppt format requires conversion."
        }

    try:
        from pptx import Presentation
        prs = Presentation(path)
        text_runs = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        text = paragraph.text.strip()
                        if text:
                            text_runs.append(text)

        full_text = "\n".join(text_runs)
        if not full_text.strip():
            return {
                "text": "",
                "status": "NO_TEXT_EXTRACTED",
                "error": None
            }

        return {
            "text": full_text,
            "status": "SUCCESS",
            "error": None
        }
    except ImportError:
        logger.error("python-pptx package is missing.")
        return {
            "text": "",
            "status": "MISSING_DEPENDENCY",
            "error": "python-pptx library not installed"
        }
    except Exception as e:
        logger.error(f"Error reading presentation {path}: {e}")
        return {
            "text": "",
            "status": "CORRUPT_FILE",
            "error": str(e)
        }
